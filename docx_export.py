"""
Build a Word (.docx) copy of the page with the suggested internal links placed in.

Legend used in the document:
  - blue underlined      : links that already exist on the page
  - green bold underlined: NEW suggested link on existing text
  - yellow highlight     : NEW suggested sentence to add (its anchor is linked)
"""
import io
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# segment = (text, href or None, is_new_link, highlighted)


def _add_link(par, text, url, new=False, highlight=False):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if new:
        rpr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1B7F3B" if new else "0563C1")
    rpr.append(color)
    if highlight:
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), "yellow")
        rpr.append(hl)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(rpr)
    r.append(t)
    h.append(r)
    par._p.append(h)


def _pattern(anchor):
    return re.compile(r"\s+".join(map(re.escape, anchor.split())), re.I)


def _link_existing(segs, anchor, url):
    """Wrap the first occurrence of anchor (in non-linked text) with the new link."""
    pat = _pattern(anchor)
    for i, (text, href, new, hl) in enumerate(segs):
        if href:
            continue
        m = pat.search(text)
        if m:
            parts = [(text[:m.start()], None, False, hl),
                     (m.group(), url, True, hl),
                     (text[m.end():], None, False, hl)]
            segs[i:i + 1] = [p for p in parts if p[0]]
            return True
    return False


def _append_sentence(segs, sentence, anchor, url):
    m = _pattern(anchor).search(sentence)
    if not m:
        return False
    segs.append((" ", None, False, False))
    parts = [(sentence[:m.start()], None, False, True),
             (m.group(), url, True, True),
             (sentence[m.end():], None, False, True)]
    segs.extend(p for p in parts if p[0])
    return True


def build_docx(title, page_url, blocks, result):
    """Return .docx bytes. `result` is the validated suggestions list from internal_links.validate()."""
    by_pid = {}
    for s in result:
        for o in s["options"]:
            by_pid.setdefault(int(o["paragraph"]), []).append((s["target_url"], o))

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title or "Page content", 0)
    p = doc.add_paragraph()
    p.add_run("Source: ").bold = True
    p.add_run(page_url)
    p = doc.add_paragraph()
    p.add_run("Legend: ").bold = True
    p.add_run("blue = existing link;  ")
    g = p.add_run("green bold = new suggested link")
    g.bold = True
    g.font.color.rgb = RGBColor(0x1B, 0x7F, 0x3B)
    p.add_run(";  ")
    p.add_run("yellow = new sentence to add").font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph()

    summary = []  # (url, anchor, type, placed)
    for b in blocks:
        if b["type"] == "heading":
            doc.add_heading(b["text"], level=min(b["level"], 9))
            continue
        par = doc.add_paragraph(style="List Bullet" if b["type"] == "li" else None)
        segs = list(b["segments"])
        for url, o in by_pid.get(b.get("pid"), []) if b.get("pid") is not None else []:
            anchor = o["anchor_text"]
            if o["type"] == "existing_text":
                ok = _link_existing(segs, anchor, url)
            else:
                ok = _append_sentence(segs, o.get("suggested_sentence", ""), anchor, url)
            summary.append((url, anchor, o["type"], ok))
        for text, href, new, hl in segs:
            if href:
                _add_link(par, text, href, new, hl)
            else:
                run = par.add_run(text)
                if hl:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    if summary:
        doc.add_page_break()
        doc.add_heading("Link suggestions summary", 1)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ("Anchor text", "Target URL", "Type", "Placed in document")):
            c.text = h
        for url, anchor, typ, ok in summary:
            cells = t.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = anchor, url, typ
            cells[3].text = "Yes" if ok else "No - add manually"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
